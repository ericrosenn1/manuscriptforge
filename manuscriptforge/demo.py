"""Offline, original synthetic corpus demonstration using the public workflow APIs."""

from __future__ import annotations

import io
import json
import stat
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document

from manuscriptforge.config import load_project_config, validate_project, write_project_config
from manuscriptforge.ingest.style_ingest import build_style_chunks, ingest_style_corpus
from manuscriptforge.models.project import ProjectConfig
from manuscriptforge.models.style import StyleProfile
from manuscriptforge.style.curation import (
    build_style_cards,
    build_style_chunk_registry,
    build_style_coverage_report,
    load_chunk_registry,
    update_style_chunk_approval,
)
from manuscriptforge.style.document_extractors import extract_style_corpus, load_extraction_manifest
from manuscriptforge.style.profiler import build_style_profile
from manuscriptforge.utils.hashing import sha256_file
from manuscriptforge.utils.io import read_json, write_json, write_text

DEMO_ID = "synthetic-distance-workflow-v1"
MANIFEST_NAME = "demo_manifest.json"
MODE = "academic_manuscript"
CORPUS_PREFIX = f"style_corpus/{MODE}"
DUPLICATE_SOURCE = f"{CORPUS_PREFIX}/methods_duplicate.docx"
REFERENCE_SENTINEL = "Synthetic reference-section sentinel: no external sources were used."

# Original release fixtures. These describe an invented protocol, not measured results.
PASSAGES = {
    "abstract": (
        "This synthetic collection describes how an imagined distance-sensor workflow could "
        "keep records interpretable as they move from collection to review. It contains no "
        "physical measurements and makes no claim about instrument performance. The passages "
        "use distinct section headings so that a reader can inspect extraction, chunk selection, "
        "and descriptive style summaries. One repeated methods passage deliberately tests "
        "whether a documented exclusion prevents duplicate text from entering the profile."
    ),
    "introduction": (
        "A short measurement record can become difficult to interpret when a processing step "
        "silently replaces its inputs. In this invented workflow, a movable target provides a "
        "simple setting for explaining that problem without using a real experiment. The "
        "documentation separates the intended collection procedure from later interpretation. "
        "This separation gives reviewers a concrete way to trace a summary back to a recorded "
        "step, while leaving the adequacy of a future physical experiment unresolved."
    ),
    "methods": (
        "We described a simulated distance sensor sampled at fixed intervals while a target "
        "moved between three marked positions. The proposed record stores the position label, "
        "sample index, and unrounded reading for each observation. A calibration offset would "
        "be applied only to a separate analysis column, leaving raw readings unchanged. Before "
        "summarizing a position, an operator would check consecutive indices and explicit "
        "units. These steps specify an imagined procedure; no sensor simulation or physical "
        "measurement was performed to produce this writing collection."
    ),
    "results": (
        "The synthetic run note illustrates how an incomplete record would be reported. A "
        "missing position label would leave that observation unresolved until its origin was "
        "checked, rather than assigning it to the nearest position. A summary table would "
        "identify the affected record and keep the original value available. This passage "
        "contains an example reporting rule, not an observed failure rate, numerical result, "
        "or assessment of a working sensor."
    ),
    "discussion": (
        "Keeping the collection record separate from its interpretation may make a future "
        "review easier to reproduce. The benefit would depend on whether labels, units, and "
        "processing choices were recorded consistently. A traceable file alone would not "
        "establish that a calibration was appropriate or that an instrument was accurate. "
        "The proposed documentation therefore treats each summary as a statement about "
        "specified inputs, with the interpretation remaining open to technical review."
    ),
    "limitations": (
        "This example is too small to represent the range of writing in a research program. "
        "Each retained section contributes one passage, and the duplicate is deliberately "
        "simple. Sentence boundaries, section labels, and linguistic indicators are evaluated "
        "with rules that can misread unfamiliar formats. The resulting profile is a "
        "descriptive view of these invented passages. It is neither a validated measure of "
        "writing quality nor evidence that a model has learned an author's preferences."
    ),
}


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(f"Demo validation failed: {message}")


def _is_link(path: Path) -> bool:
    """Include Windows junctions without requiring Python 3.12's is_junction()."""
    info = path.lstat()
    return path.is_symlink() or bool(
        getattr(info, "st_file_attributes", 0) & stat.FILE_ATTRIBUTE_REPARSE_POINT
    )


def _inventory(root: Path) -> dict[str, dict[str, Any]]:
    files: dict[str, dict[str, Any]] = {}
    pending = [root]
    while pending:
        for path in sorted(pending.pop().iterdir()):
            _require(not _is_link(path), "output contains a symbolic link or junction")
            if path.is_dir():
                pending.append(path)
            elif path.is_file() and path != root / MANIFEST_NAME:
                files[path.relative_to(root).as_posix()] = {
                    "sha256": sha256_file(path), "bytes": path.stat().st_size
                }
            elif path != root / MANIFEST_NAME:
                raise ValueError("Demo output contains an unsupported filesystem entry")
    return dict(sorted(files.items()))


def _output_path(project: Path, relative: str) -> Path:
    path = project / relative
    _require(not Path(relative).is_absolute() and path.resolve().is_relative_to(project.resolve()), "output path leaves demo directory")
    return path


def _write_documents(project: Path) -> None:
    def document_text(sections: list[str], markdown: bool) -> str:
        prefix = "## " if markdown else ""
        return "\n\n".join(
            f"{prefix}{section.title()}\n\n{PASSAGES[section]}" for section in sections
        ) + "\n"

    write_text(
        project / CORPUS_PREFIX / "protocol.md",
        document_text(["abstract", "introduction", "methods"], True)
        + f"\n## References\n\n{REFERENCE_SENTINEL}\n",
    )
    write_text(
        project / CORPUS_PREFIX / "run_note.txt",
        document_text(["results", "discussion", "limitations"], False),
    )
    document = Document()
    properties = document.core_properties
    properties.author = "ManuscriptForge synthetic example"
    properties.last_modified_by = ""
    properties.title = "Synthetic methods duplicate"
    properties.subject = "Original nonclinical demonstration data"
    properties.comments = "Generated from original synthetic text by manuscriptforge.demo."
    properties.keywords = "synthetic, demonstration"
    properties.created = properties.modified = datetime(2000, 1, 1, tzinfo=UTC)
    document.add_heading("Methods", level=1)
    document.add_paragraph(PASSAGES["methods"])
    raw = io.BytesIO()
    document.save(raw)
    # python-docx uses current ZIP timestamps; normalize only this new demo asset.
    with ZipFile(raw) as source, ZipFile(project / DUPLICATE_SOURCE, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            entry.date_time = (2000, 1, 1, 0, 0, 0)
            target.writestr(entry, source.read(entry.filename))


def _validate_outputs(project: Path) -> dict[str, Any]:
    config = load_project_config(project)
    _require(config.privacy.local_only and config.llm.provider == "mock", "offline configuration")
    _require(config.style.curation.require_chunk_approval, "strict chunk approval is required")
    extractions = load_extraction_manifest(project)
    _require(len(extractions) == 3, "expected three source documents")
    _require({item.relative_path for item in extractions} == {
        f"{CORPUS_PREFIX}/protocol.md", f"{CORPUS_PREFIX}/run_note.txt", DUPLICATE_SOURCE,
    }, "unexpected source paths")
    _require({item.extension for item in extractions} == {".md", ".txt", ".docx"}, "source formats")
    for item in extractions:
        _require(item.extraction_status == "extracted", "every source must extract successfully")
        _require(item.content_hash == sha256_file(project / item.relative_path), "source hash changed")
        _require(bool(item.extracted_text_path), "missing extracted text path")
        _require(bool(_output_path(project, str(item.extracted_text_path)).read_text(encoding="utf-8").strip()), "empty extracted text")
    rows = load_chunk_registry(project)
    _require(len(rows) == 7 and not any(row.missing for row in rows), "expected seven present chunks")
    approved = [row for row in rows if row.approval_status == "approved"]
    excluded = [row for row in rows if row.approval_status == "excluded"]
    _require(len(approved) == 6 and len(excluded) == 1, "expected six approved and one excluded chunk")
    _require(excluded[0].source_file == DUPLICATE_SOURCE, "the DOCX duplicate must be excluded")
    _require(all("style_profile" in row.approved_for for row in approved), "approved profile use")
    repeated = [row for row in rows if row.content_hash == excluded[0].content_hash]
    _require(len(repeated) == 2, "the exact duplicate must match one retained chunk")
    _require(all("repeated_boilerplate" in row.warning_flags for row in repeated), "duplicate warning flags")
    profile = StyleProfile.model_validate(read_json(project / "outputs/profile/style_profile.json"))
    features = profile.global_features
    _require(set(features["chunk_ids"]) == {row.chunk_id for row in approved}, "profile and approval IDs differ")
    _require(DUPLICATE_SOURCE not in profile.corpus_files, "duplicate entered profile")
    _require(features["word_count"] > 300 and features["sentence_count"] > 12, "profile lacks substantive text")
    _require(features["section_distribution"] == dict.fromkeys(PASSAGES, 1), "expected six represented sections")
    _require(REFERENCE_SENTINEL not in json.dumps(profile.model_dump()), "References entered profile")
    for section, passage in PASSAGES.items():
        _require(profile.section_features[section]["chunk_count"] == 1, f"{section} section count")
        _require(profile.section_features[section]["representative_examples"][0]["text"] == passage, f"{section} text changed")
    coverage = read_json(project / "planning/intake/style_coverage_report.json")
    coverage_rows = {row["section_type"]: row for row in coverage["coverage"]}
    _require(all(coverage_rows[section]["approved_chunks"] == 1 for section in PASSAGES), "coverage counts")
    _require(all(coverage_rows[section]["coverage_strength"] == "sparse" for section in PASSAGES), "coverage labels")
    cards = read_json(project / "planning/intake/style_cards/style_cards_manifest.json")
    _require(len(cards["cards"]) == 8, "expected eight style cards")
    _require(all(_output_path(project, card["path"]).stat().st_size > 0 for card in cards["cards"]), "empty style card")
    validation = read_json(project / "project_validation.json")
    _require(not validation["errors"], "project configuration has errors")
    return {
        "demo_id": DEMO_ID,
        "synthetic": True,
        "validation_status": "PASS",
        "source_documents": len(extractions),
        "input_formats": sorted(item.extension for item in extractions),
        "extracted_chunks": len(rows),
        "approved_chunks": len(approved),
        "excluded_duplicate_chunks": len(excluded),
        "profile_source_documents": len(profile.corpus_files),
        "profile_word_count": features["word_count"],
        "profile_sentence_count": features["sentence_count"],
        "section_counts": dict(sorted(Counter(row.section_type for row in approved).items())),
        "style_cards": len(cards["cards"]),
        "project_validation_warnings": validation["warnings"],
        "checks": [
            "source_hashes_preserved", "nonempty_extraction", "references_excluded",
            "exact_duplicate_flagged", "duplicate_explicitly_excluded", "approval_decisions_preserved",
            "profile_matches_approved_chunks", "section_text_preserved", "coverage_and_cards_validated",
        ],
        "reports": {
            "walkthrough": "DEMO_REPORT.md",
            "profile": "outputs/profile/style_profile.json",
            "style_guide": "outputs/profile/style_guide.md",
            "chunk_review": "planning/intake/style_chunks/style_chunk_review.md",
            "coverage": "planning/intake/style_coverage_report.md",
            "cards": "planning/intake/style_cards/style_cards_manifest.json",
        },
    }


def run_demo(output_dir: Path | str) -> dict[str, Any]:
    """Build and validate a local synthetic demo, or verify an intact previous run.

    The destination must be absent, empty, or an unchanged completed demo. Existing
    completed runs are checked without rewriting files. No model or network is used.
    """
    project = Path(output_dir).absolute()
    for parent in [project, *project.parents]:
        if parent.exists() or parent.is_symlink():
            _require(not _is_link(parent), "destination traverses a symbolic link or junction")
    if project.exists():
        _require(project.is_dir(), "destination must be a directory")
        if any(project.iterdir()):
            marker = project / MANIFEST_NAME
            _require(marker.is_file(), "destination is nonempty and is not a completed demo; choose a new directory")
            _require(not _is_link(marker), "demo manifest is a symbolic link or junction")
            manifest = read_json(marker)
            _require(manifest.get("demo_id") == DEMO_ID, "destination belongs to another demo")
            # Reject unrelated directories before traversing their contents.
            actual = _inventory(project)
            _require(manifest.get("files") == actual, "existing demo files changed; choose a new directory")
            summary = _validate_outputs(project)
            _require(read_json(project / "demo_summary.json") == summary, "saved summary differs from content")
            return summary
    project.mkdir(parents=True, exist_ok=True)
    config = ProjectConfig(project_name="Synthetic distance workflow", field="measurement documentation",
                           target_audience="technical workflow reviewers")
    config.privacy.local_only = True
    config.style.curation.require_chunk_approval = True
    config.style.curation.include_needs_review_chunks = False
    config.style.extraction.fail_on_extraction_error = True
    write_project_config(project, config)
    for folder in ["inputs/results_tables", "inputs/source_pdfs", "outputs/profile"]:
        (project / folder).mkdir(parents=True, exist_ok=True)
    for section in ["abstract", "methods"]:
        write_text(project / "inputs" / f"{section}.md", PASSAGES[section] + "\n")
    _write_documents(project)
    validation = validate_project(project)
    _require(validation.ok, "; ".join(validation.errors))
    write_json(project / "project_validation.json", {
        "errors": validation.errors, "warnings": validation.warnings, "infos": validation.infos,
    })
    source_hashes = {path.name: sha256_file(path) for path in (project / CORPUS_PREFIX).iterdir()}
    extract_style_corpus(project, mode=MODE)
    samples = ingest_style_corpus(project)
    chunks = build_style_chunks(project, samples)
    _require(len(chunks) == 7 and all(REFERENCE_SENTINEL not in chunk.text for chunk in chunks), "extraction and reference exclusion")
    registry = build_style_chunk_registry(project, mode=MODE)
    for source in sorted({row.source_file for row in registry.rows}):
        duplicate = source == DUPLICATE_SOURCE
        update_style_chunk_approval(
            project, source_file=source, approve=not duplicate, exclude=duplicate,
            approve_for="style_profile" if not duplicate else None,
            note="Scripted demo decision: exact duplicate excluded." if duplicate else
                 "Scripted demo decision: original synthetic passage retained for descriptive profiling.",
        )
    decisions = {row.chunk_id: (row.approval_status, row.approved_for, row.notes)
                 for row in load_chunk_registry(project)}
    rebuilt = build_style_chunk_registry(project, mode=MODE)
    _require(decisions == {row.chunk_id: (row.approval_status, row.approved_for, row.notes)
                          for row in rebuilt.rows}, "registry rebuild lost approval decisions")
    build_style_profile(project, run_dir=project / "outputs/profile", samples=samples)
    build_style_coverage_report(project, mode=MODE)
    build_style_cards(project, mode=MODE)
    _require(source_hashes == {path.name: sha256_file(path) for path in (project / CORPUS_PREFIX).iterdir()}, "source documents changed")
    summary = _validate_outputs(project)
    write_json(project / "demo_summary.json", summary)
    write_text(project / "DEMO_REPORT.md", (
        "# Synthetic corpus demonstration\n\n"
        "This original, nonclinical example describes an invented measurement protocol. "
        "The demo performs no sensor experiment or training and calls no model provider.\n\n"
        f"Three documents in Markdown, text, and DOCX produced {summary['extracted_chunks']} chunks. "
        "Scripted curation retained six passages and explicitly excluded the DOCX methods duplicate. "
        "Both copies carry a repeated_boilerplate warning; that warning alone does not remove text. "
        "The registry was rebuilt and its decisions checked before profiling.\n\n"
        f"The approved profile contains {summary['profile_word_count']} words and "
        f"{summary['profile_sentence_count']} sentences across six sections. Each section has one "
        "approved chunk, so coverage is sparse. These are descriptive counts, not quality scores.\n\n"
        "Start with [the style guide](outputs/profile/style_guide.md), "
        "[chunk review](planning/intake/style_chunks/style_chunk_review.md), and "
        "[coverage report](planning/intake/style_coverage_report.md). "
        "[demo_summary.json](demo_summary.json) records the validated relationships using relative paths.\n\n"
        "The general project validator also reports missing optional drafting materials, tables, "
        "and citations. They are intentionally absent from this corpus-only demonstration. "
        "Local extraction records include source paths and text. Review them before sharing.\n\n"
        "Run the same demo command again to verify all saved file hashes and output relationships. "
        "An intact completed run is reused without changes. Modified or unrelated output directories "
        "are rejected; select a new directory to create another run.\n"
    ))
    write_json(project / MANIFEST_NAME, {"demo_id": DEMO_ID, "files": _inventory(project)})
    return summary
