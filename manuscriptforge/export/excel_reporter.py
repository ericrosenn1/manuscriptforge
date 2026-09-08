from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from pydantic import BaseModel

from manuscriptforge.models.audit import AuditFinding
from manuscriptforge.models.citation import CitationRecord
from manuscriptforge.models.claim import ScientificClaim


def _plain(value: Any) -> Any:
    if isinstance(value, BaseModel):
        return value.model_dump(mode="json")
    if isinstance(value, list):
        return [_plain(item) for item in value]
    if isinstance(value, dict):
        return {key: _plain(item) for key, item in value.items()}
    return value


def _cell(value: Any) -> Any:
    value = _plain(value)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def write_rows_xlsx(path: Path, rows: list[dict[str, Any]], sheet_name: str = "Sheet1") -> None:
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31] or "Sheet1"
    if rows:
        headers = list(rows[0].keys())
    else:
        headers = ["message"]
        rows = [{"message": "No records"}]
    worksheet.append(headers)
    for row in rows:
        worksheet.append([_cell(row.get(header)) for header in headers])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 60)
    workbook.save(path)


def write_workbook(path: Path, sheets: dict[str, list[dict[str, Any]]]) -> None:
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    first = True
    for sheet_name, rows in sheets.items():
        worksheet = workbook.active if first else workbook.create_sheet()
        first = False
        worksheet.title = sheet_name[:31] or "Sheet"
        if rows:
            headers = list(rows[0].keys())
        else:
            headers = ["message"]
            rows = [{"message": "No records"}]
        worksheet.append(headers)
        for row in rows:
            worksheet.append([_cell(row.get(header)) for header in headers])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_cells in worksheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(
                max(max_length + 2, 12), 60
            )
    workbook.save(path)


def write_claim_registry_xlsx(path: Path, claims: list[ScientificClaim]) -> None:
    rows = []
    evidence_rows = []
    for claim in claims:
        rows.append(
            {
                "claim_id": claim.claim_id,
                "claim_text": claim.claim_text,
                "section_target": claim.section_target,
                "support_strength": claim.support_strength,
                "claim_type": claim.claim_type,
                "evidence_type": claim.evidence_type,
                "evidence_ids": ", ".join(item.evidence_id for item in claim.evidence_items),
                "allowed_language": ", ".join(claim.allowed_language),
                "forbidden_language": ", ".join(claim.forbidden_language),
                "citation_ids": ", ".join(claim.citation_ids),
                "needs_citation": claim.needs_citation,
                "needs_human_review": claim.needs_human_review,
                "notes": claim.notes,
            }
        )
        for evidence in claim.evidence_items:
            evidence_rows.append(
                {
                    "claim_id": claim.claim_id,
                    "evidence_id": evidence.evidence_id,
                    "evidence_type": evidence.evidence_type,
                    "source_label": evidence.source_label,
                    "source_path": evidence.source_path,
                    "locator": evidence.locator,
                    "text": evidence.text,
                    "content_hash": evidence.content_hash,
                }
            )
    write_workbook(path, {"claims": rows, "evidence": evidence_rows})


def write_citation_audit_xlsx(
    path: Path, findings: list[AuditFinding], citations: list[CitationRecord] | None = None
) -> None:
    rows = [finding.model_dump(mode="json") for finding in findings]
    if not rows and citations:
        rows = [
            {
                "finding_id": f"cit_ok_{index}",
                "severity": "info",
                "category": "citation_registry",
                "message": f"Registered citation: {citation.title}",
                "location": citation.citation_id,
                "suggested_fix": "",
                "related_claim_ids": [],
                "related_citation_ids": [citation.citation_id],
            }
            for index, citation in enumerate(citations, start=1)
        ]
    citation_rows = [
        {
            "citation_id": citation.citation_id,
            "title": citation.title,
            "authors": "; ".join(citation.authors),
            "year": citation.year,
            "doi": citation.doi,
            "pmid": citation.pmid,
            "journal": citation.journal,
            "source_path": citation.source_path,
            "notes": citation.notes,
        }
        for citation in citations or []
    ]
    write_workbook(path, {"citation_audit": rows, "citation_registry": citation_rows})


def write_findings_xlsx(path: Path, findings: list[AuditFinding], sheet_name: str = "findings") -> None:
    write_rows_xlsx(path, [finding.model_dump(mode="json") for finding in findings], sheet_name)


def _claim_rows(claims: list[ScientificClaim]) -> list[dict[str, Any]]:
    return [
        {
            "claim_id": claim.claim_id,
            "claim_text": claim.claim_text,
            "section_target": claim.section_target,
            "support_strength": claim.support_strength,
            "claim_type": claim.claim_type,
            "evidence_ids": ", ".join(item.evidence_id for item in claim.evidence_items),
            "citation_ids": ", ".join(claim.citation_ids),
            "needs_citation": claim.needs_citation,
            "needs_human_review": claim.needs_human_review,
            "notes": claim.notes,
        }
        for claim in claims
    ]


def _citation_rows(citations: list[CitationRecord]) -> list[dict[str, Any]]:
    return [
        {
            "citation_id": citation.citation_id,
            "title": citation.title,
            "authors": "; ".join(citation.authors),
            "year": citation.year,
            "doi": citation.doi,
            "pmid": citation.pmid,
            "pmcid": citation.pmcid,
            "arxiv_id": citation.arxiv_id,
            "url": citation.url,
            "journal": citation.journal,
            "metadata_status": citation.metadata_status,
            "metadata_confidence": citation.metadata_confidence,
            "source_kind": citation.source_kind,
            "retrieval_notes": citation.retrieval_notes,
        }
        for citation in citations
    ]


def _table_rows(tables: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "table_id": table.get("table_id"),
            "source_label": table.get("source_label"),
            "source_path": table.get("source_path"),
            "description": (table.get("schema") or {}).get("description") if table.get("schema") else None,
            "row_count": table.get("row_count"),
            "column_count": table.get("column_count") or len(table.get("columns", [])),
            "numeric_columns": ", ".join(table.get("numeric_columns", [])),
            "p_value_columns": ", ".join(table.get("p_value_columns", [])),
            "adjusted_p_value_columns": ", ".join(table.get("adjusted_p_value_columns", [])),
            "effect_size_columns": ", ".join(table.get("effect_size_columns", [])),
            "warnings": "; ".join(table.get("warnings", [])),
        }
        for table in tables
    ]


def _source_map_rows(source_map: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for section in source_map.get("sections", []):
        for paragraph in section.get("paragraphs", []):
            rows.append(
                {
                    "section_name": section.get("section_name"),
                    "paragraph_index": paragraph.get("paragraph_index"),
                    "claim_ids": ", ".join(paragraph.get("claim_ids", [])),
                    "citation_ids": ", ".join(paragraph.get("citation_ids", [])),
                    "style_mode": paragraph.get("style_mode", section.get("style_mode")),
                    "style_chunk_ids": ", ".join(paragraph.get("style_chunk_ids", [])),
                    "unresolved_flags": "; ".join(paragraph.get("unresolved_flags", [])),
                    "generated_by": paragraph.get("generated_by"),
                    "text": paragraph.get("text"),
                }
            )
    return rows


def _manifest_rows(manifest: dict[str, Any]) -> list[dict[str, Any]]:
    if not manifest:
        return []
    rows = [
        {"field": "command", "value": manifest.get("command")},
        {"field": "created_utc", "value": manifest.get("created_utc")},
        {"field": "completed_utc", "value": manifest.get("completed_utc")},
        {"field": "config_hash", "value": manifest.get("config_hash")},
        {"field": "files", "value": len(manifest.get("files", []))},
        {"field": "input_hashes", "value": len(manifest.get("input_hashes", {}))},
        {"field": "output_hashes", "value": len(manifest.get("output_hashes", {}))},
    ]
    for key, value in manifest.get("extra", {}).items():
        rows.append({"field": f"extra.{key}", "value": value})
    return rows


def write_audit_workbook(
    path: Path,
    *,
    claims: list[ScientificClaim],
    citations: list[CitationRecord],
    findings: list[AuditFinding],
    revision_questions: list[dict[str, Any]],
    tables: list[dict[str, Any]],
    source_map: dict[str, Any],
    manifest: dict[str, Any],
) -> None:
    write_workbook(
        path,
        {
            "Claims": _claim_rows(claims),
            "Citations": _citation_rows(citations),
            "Audit Findings": [finding.model_dump(mode="json") for finding in findings],
            "Revision Questions": revision_questions,
            "Table Inventory": _table_rows(tables),
            "Source Map": _source_map_rows(source_map),
            "Run Manifest Summary": _manifest_rows(manifest),
        },
    )
