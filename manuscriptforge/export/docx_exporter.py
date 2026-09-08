from __future__ import annotations

from pathlib import Path

from manuscriptforge.models.manuscript import Manuscript


def export_docx(path: Path, manuscript: Manuscript) -> None:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError("python-docx is required for DOCX export") from exc

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(manuscript.title, level=0)
    if manuscript.title_candidates:
        doc.add_heading("Title Candidates", level=1)
        for candidate in manuscript.title_candidates:
            doc.add_paragraph(candidate, style="List Bullet")
    if manuscript.abstract:
        doc.add_heading("Abstract", level=1)
        doc.add_paragraph(manuscript.abstract)
    for section in manuscript.sections:
        if section.section_name.lower() == "abstract":
            continue
        doc.add_heading(section.section_name, level=1)
        for paragraph in section.content.split("\n\n"):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())
        if section.claim_ids or section.citation_ids:
            doc.add_paragraph(
                "Traceability: "
                + f"claims={', '.join(section.claim_ids) or 'none'}; "
                + f"citations={', '.join(section.citation_ids) or 'none'}"
            )
        for flag in section.unresolved_flags:
            doc.add_paragraph(f"[UNRESOLVED: {flag}]")
        for note in section.revision_notes:
            doc.add_paragraph(f"[REVISION NOTE: {note}]")
    if manuscript.tables:
        doc.add_heading("Table Inventory", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        headers = ["Table", "Rows", "Columns", "Source"]
        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header
        for table_summary in manuscript.tables:
            row = table.add_row().cells
            row[0].text = str(table_summary.get("source_label", "table"))
            row[1].text = str(table_summary.get("row_count", 0))
            row[2].text = ", ".join(table_summary.get("columns", []))
            row[3].text = str(table_summary.get("source_path", ""))
    if manuscript.figure_legends:
        doc.add_heading("Figure Legends", level=1)
        for legend in manuscript.figure_legends:
            doc.add_paragraph(legend, style="List Bullet")
    if manuscript.references:
        doc.add_heading("References", level=1)
        for citation in manuscript.references:
            authors = ", ".join(citation.authors) if citation.authors else "Unknown authors"
            year = citation.year or "n.d."
            doc.add_paragraph(f"{authors} ({year}). {citation.title}. {citation.journal or ''}")
    if manuscript.ai_disclosure:
        doc.add_heading("AI-Use Disclosure Draft", level=1)
        doc.add_paragraph(manuscript.ai_disclosure)
    doc.add_heading("Review Appendix", level=1)
    doc.add_paragraph("Unsupported claims, citation gaps, and unresolved flags should be resolved by the human author before submission.")
    for section in manuscript.sections:
        if section.unresolved_flags or section.revision_notes:
            doc.add_heading(section.section_name, level=2)
            for flag in section.unresolved_flags:
                doc.add_paragraph(f"Unresolved: {flag}", style="List Bullet")
            for note in section.revision_notes:
                doc.add_paragraph(f"Revision note: {note}", style="List Bullet")
    doc.save(str(path))
