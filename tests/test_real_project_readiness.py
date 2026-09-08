from __future__ import annotations

from pathlib import Path
from shutil import copytree

import yaml

from manuscriptforge.audit.journal_audit import audit_journal_profile
from manuscriptforge.audit.no_invention_audit import audit_no_invention
from manuscriptforge.drafting.revision_questions import generate_revision_question_records
from manuscriptforge.export.docx_exporter import export_docx
from manuscriptforge.export.excel_reporter import write_audit_workbook
from manuscriptforge.export.latex_exporter import manuscript_to_latex
from manuscriptforge.ingest.bibtex_ingest import ingest_bibtex
from manuscriptforge.ingest.metadata_enrichment import MockMetadataProvider, enrich_citations
from manuscriptforge.ingest.pdf_ingest import summarize_pdf
from manuscriptforge.ingest.project_ingest import ingest_project
from manuscriptforge.models.audit import AuditFinding
from manuscriptforge.models.citation import CitationRecord
from manuscriptforge.models.claim import EvidenceItem, ScientificClaim
from manuscriptforge.models.manuscript import Manuscript, ManuscriptSection
from manuscriptforge.pipeline.diffing import run_diff
from manuscriptforge.pipeline.workflow import run_ingest
from manuscriptforge.utils.io import create_run_dir, read_json, write_json


def _copy_example(tmp_path: Path) -> Path:
    project_dir = tmp_path / "project_minimal"
    copytree(Path("tests/fixtures/synthetic_project"), project_dir)
    return project_dir


def _config(project_dir: Path) -> dict:
    return yaml.safe_load((project_dir / "project.yaml").read_text(encoding="utf-8"))


def _write_config(project_dir: Path, data: dict) -> None:
    (project_dir / "project.yaml").write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")


def test_source_list_parsing_from_txt_md_json_bib_and_ris(tmp_path) -> None:
    project_dir = _copy_example(tmp_path)
    (project_dir / "inputs" / "references.ris").write_text(
        "TY  - JOUR\nTI  - RIS Demo Source\nAU  - Example, Erin\nPY  - 2023\nDO  - 10.1000/risdemo\nER  -\n",
        encoding="utf-8",
    )
    (project_dir / "inputs" / "references.txt").write_text(
        "1. 2022 Plain text source. Journal Name. doi:10.1000/textdemo\nPMID: 12345678\n",
        encoding="utf-8",
    )
    (project_dir / "inputs" / "source_list.md").write_text(
        "- Manual markdown source 2021 with https://example.org/source\n",
        encoding="utf-8",
    )
    records = ingest_bibtex(project_dir)
    kinds = {record.source_kind for record in records}
    assert {"bibtex", "ris", "plain_text", "pmid_list", "manual_list"}.issubset(kinds)
    assert any(record.doi == "10.1000/risdemo" for record in records)
    assert any(record.pmid == "12345678" for record in records)


def test_metadata_enrichment_offline_skipped_and_mock_success(tmp_path) -> None:
    project_dir = _copy_example(tmp_path)
    config = _config(project_dir)
    config["sources"]["enrich_doi"] = True
    _write_config(project_dir, config)
    run_dir = run_ingest(project_dir, enrich_sources=True)
    citations = read_json(run_dir / "citation_registry.json")
    assert any(citation["metadata_status"] == "offline_skipped" for citation in citations)
    manifest = read_json(run_dir / "run_manifest.json")
    assert manifest["metadata_cache_hashes"]

    citation = CitationRecord(citation_id="cit_x", title="Old", doi="10.1000/mock")
    config_model = __import__("manuscriptforge.config", fromlist=["load_project_config"]).load_project_config(
        project_dir
    )
    enriched, cache_files = enrich_citations(
        [citation],
        project_dir,
        config_model,
        provider=MockMetadataProvider({"10.1000/mock": {"title": "Mock Enriched Title", "year": "2026"}}),
        force=True,
    )
    assert enriched[0].metadata_status == "enriched"
    assert enriched[0].title == "Mock Enriched Title"
    assert cache_files


def test_scanned_pdf_heuristic_with_monkeypatched_extractor(tmp_path, monkeypatch) -> None:
    project_dir = _copy_example(tmp_path)
    pdf_path = project_dir / "inputs" / "source_pdfs" / "scan.pdf"
    pdf_path.write_bytes(b"%PDF fake")
    import manuscriptforge.ingest.pdf_ingest as pdf_ingest

    monkeypatch.setattr(pdf_ingest, "extract_pdf_pages", lambda path: ["tiny", ""])
    summary = summarize_pdf(pdf_path, project_dir)
    assert summary.extraction_status == "likely_scanned"
    assert summary.needs_ocr is True


def test_table_schema_hints_and_missing_schema_column_warning(tmp_path) -> None:
    project_dir = _copy_example(tmp_path)
    config = _config(project_dir)
    config["tables"]["schemas"]["results_summary.csv"]["column_roles"]["missing_column"] = "p_value"
    _write_config(project_dir, config)
    result = __import__("manuscriptforge.config", fromlist=["validate_project"]).validate_project(project_dir)
    assert any("missing_column" in warning for warning in result.warnings)

    run_dir = create_run_dir(project_dir)
    ingested = ingest_project(project_dir, run_dir)
    table = ingested["tables"][0]
    assert table["role_columns"]["effect_size"] == ["log2_fold_change"]
    assert table["threshold_counts"]


def test_journal_audit_flags_required_items() -> None:
    manuscript = Manuscript(
        title="Short",
        abstract=" ".join(["word"] * 310),
        sections=[ManuscriptSection(section_name="Results", section_type="results", content="Results.")],
    )
    from manuscriptforge.models.project import ProjectConfig

    findings, report = audit_journal_profile(manuscript, ProjectConfig(target_journal="generic_biomedical"))
    assert any("Required section missing" in finding.message for finding in findings)
    assert any("Abstract has" in finding.message for finding in findings)
    assert "Journal/Profile Audit" in report


def test_run_diff_detects_text_claim_and_audit_changes(tmp_path) -> None:
    project_dir = tmp_path / "project"
    run_a = project_dir / "outputs" / "runs" / "20260101_000000"
    run_b = project_dir / "outputs" / "runs" / "20260102_000000"
    run_a.mkdir(parents=True)
    run_b.mkdir(parents=True)
    (run_a / "manuscript.md").write_text("old text\n", encoding="utf-8")
    (run_b / "manuscript.md").write_text("new text\n", encoding="utf-8")
    write_json(run_a / "claim_registry.json", [])
    write_json(run_b / "claim_registry.json", [{"claim_id": "clm_added", "support_strength": "direct"}])
    write_json(run_a / "citation_registry.json", [])
    write_json(run_b / "citation_registry.json", [])
    write_json(run_a / "audit_findings.json", [{"finding_id": "aud_old", "severity": "serious"}])
    write_json(run_b / "audit_findings.json", [])
    write_json(run_a / "run_manifest.json", {"config_hash": "a", "files": []})
    write_json(run_b / "run_manifest.json", {"config_hash": "b", "files": []})
    write_json(run_a / "style_profile.json", {"corpus_files": []})
    write_json(run_b / "style_profile.json", {"corpus_files": []})

    diff_dir = run_diff(project_dir)
    report = (diff_dir / "diff_report.md").read_text(encoding="utf-8")
    assert "clm_added" in report
    assert "Resolved serious findings: 1" in report
    assert "-old text" in report


def test_no_invention_flags_unknown_numeric_and_citation_but_allows_known() -> None:
    evidence = EvidenceItem(
        evidence_id="ev_1",
        evidence_type="table",
        source_path="x.csv",
        source_label="x",
        text="value 1.2",
        metadata={},
        content_hash="abc",
    )
    claim = ScientificClaim(
        claim_id="clm_1",
        claim_text="The value was 1.2.",
        normalized_claim="the value was 1.2.",
        section_target="Results",
        evidence_items=[evidence],
        support_strength="direct",
        claim_type="result",
    )
    citation = CitationRecord(citation_id="cit_1", title="Known", year="2024", doi="10.1000/known")
    clean = Manuscript(
        title="Clean",
        sections=[
            ManuscriptSection(
                section_name="Results",
                section_type="results",
                content="The value was 1.2 [cit_1].",
                claim_ids=["clm_1"],
                citation_ids=["cit_1"],
            )
        ],
    )
    clean_findings, _ = audit_no_invention(clean, [claim], [citation])
    assert not clean_findings

    dirty = Manuscript(
        title="Dirty",
        sections=[
            ManuscriptSection(
                section_name="Results",
                section_type="results",
                content="The value was 999 [cit_missing]. doi:10.1000/fake",
                claim_ids=["clm_missing"],
                citation_ids=["cit_missing"],
            )
        ],
    )
    findings, report = audit_no_invention(dirty, [claim], [citation])
    assert any("Numeric value" in finding.message for finding in findings)
    assert any("citation ID" in finding.message for finding in findings)
    assert "No-Invention Report" in report


def test_revision_question_categories_are_structured_and_deduped() -> None:
    findings = [
        AuditFinding(
            finding_id="aud_repro",
            severity="warning",
            category="reproducibility",
            message="Methods may be missing software versions.",
            location="Methods",
        ),
        AuditFinding(
            finding_id="aud_journal",
            severity="warning",
            category="journal_compliance",
            message="Required statement may be missing: data availability.",
            location="statements",
        ),
    ]
    claim = ScientificClaim(
        claim_id="clm_unsupported",
        claim_text="The marker causes response.",
        normalized_claim="the marker causes response.",
        section_target="Discussion",
        support_strength="unsupported",
        claim_type="interpretation",
        needs_human_review=True,
    )
    questions = generate_revision_question_records([claim], findings)
    categories = {question["category"] for question in questions}
    assert {"Claim strength", "Code/software reproducibility", "Journal compliance"}.issubset(categories)
    assert len({question["question_id"] for question in questions}) == len(questions)


def test_richer_exports_include_expected_appendices_and_sheets(tmp_path) -> None:
    from docx import Document
    from openpyxl import load_workbook

    manuscript = Manuscript(
        title="Special & Title_1",
        abstract="Abstract.",
        sections=[
            ManuscriptSection(
                section_name="Results",
                section_type="results",
                content="A result.",
                claim_ids=["clm_1"],
                citation_ids=["cit_1"],
                unresolved_flags=["citation needed"],
                revision_notes=["Style target: concise."],
            )
        ],
    )
    md = __import__("manuscriptforge.export.markdown_exporter", fromlist=["manuscript_to_markdown"]).manuscript_to_markdown(
        manuscript
    )
    assert "## Source Traceability Appendix" in md
    assert "## Unresolved Flags Appendix" in md

    docx_path = tmp_path / "manuscript.docx"
    export_docx(docx_path, manuscript)
    doc = Document(docx_path)
    assert any("Review Appendix" in paragraph.text for paragraph in doc.paragraphs)

    latex = manuscript_to_latex(manuscript)
    assert r"Special \& Title\_1" in latex

    workbook_path = tmp_path / "audit_workbook.xlsx"
    write_audit_workbook(
        workbook_path,
        claims=[],
        citations=[],
        findings=[],
        revision_questions=[],
        tables=[],
        source_map={"sections": []},
        manifest={"command": "test", "files": [], "input_hashes": {}, "output_hashes": {}, "extra": {}},
    )
    workbook = load_workbook(workbook_path)
    assert {
        "Claims",
        "Citations",
        "Audit Findings",
        "Revision Questions",
        "Table Inventory",
        "Source Map",
        "Run Manifest Summary",
    }.issubset(set(workbook.sheetnames))
