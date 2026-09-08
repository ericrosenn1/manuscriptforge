
from manuscriptforge.export.docx_exporter import export_docx
from manuscriptforge.export.excel_reporter import write_claim_registry_xlsx
from manuscriptforge.export.markdown_exporter import export_markdown
from manuscriptforge.models.claim import EvidenceItem, ScientificClaim
from manuscriptforge.models.manuscript import Manuscript, ManuscriptSection


def test_markdown_export_writes_manuscript(tmp_path) -> None:
    manuscript = Manuscript(
        title="Example",
        abstract="Short abstract.",
        title_candidates=["Example", "Alternative"],
        sections=[
            ManuscriptSection(
                section_name="Results",
                section_type="results",
                content="A result.",
                claim_ids=["clm_1"],
                citation_ids=["cit_1"],
            )
        ],
        tables=[{"source_label": "results", "row_count": 1, "columns": ["marker"], "source_path": "x.csv"}],
    )
    out = tmp_path / "manuscript.md"
    export_markdown(out, manuscript)
    text = out.read_text(encoding="utf-8")
    assert "# Example" in text
    assert "### Traceability" in text
    assert "## Table Inventory" in text


def test_docx_export_writes_file(tmp_path) -> None:
    manuscript = Manuscript(
        title="Example",
        sections=[ManuscriptSection(section_name="Methods", section_type="methods", content="A method.")],
    )
    out = tmp_path / "manuscript.docx"
    export_docx(out, manuscript)
    assert out.exists()
    assert out.stat().st_size > 0


def test_claim_registry_xlsx_has_evidence_sheet(tmp_path) -> None:
    from openpyxl import load_workbook

    evidence = EvidenceItem(
        evidence_id="ev_1",
        evidence_type="table",
        source_path="inputs/results_tables/a.csv",
        source_label="a",
        text="A table row.",
        metadata={},
        content_hash="abc",
    )
    claim = ScientificClaim(
        claim_id="clm_1",
        claim_text="A result was reported.",
        normalized_claim="a result was reported.",
        section_target="Results",
        evidence_items=[evidence],
        support_strength="direct",
        claim_type="result",
    )
    out = tmp_path / "claims.xlsx"
    write_claim_registry_xlsx(out, [claim])
    workbook = load_workbook(out)
    assert {"claims", "evidence"}.issubset(set(workbook.sheetnames))
